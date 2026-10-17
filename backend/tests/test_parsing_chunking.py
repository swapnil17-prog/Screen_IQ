import io

import pytest
from docx import Document

from app.llm.mock import mock_resume
from app.services.chunking import build_resume_chunks, build_section_texts, split_text
from app.services.parsing import ParseError, parse_resume
from tests.conftest import read_sample


def _docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestParsing:
    def test_docx_paragraphs_and_tables(self):
        data = _docx(["Jane Doe", "Senior Engineer with ten years of experience building distributed systems."],
                     table=[["Skills", "Python, Go, Kubernetes"]])
        text = parse_resume(data, "cv.docx")
        assert "Jane Doe" in text
        assert "Python, Go, Kubernetes" in text

    def test_txt(self):
        text = parse_resume(read_sample("resume_ada_strong.txt").encode(), "resume.txt")
        assert text.startswith("Ada Okafor")

    def test_legacy_doc_rejected_with_guidance(self):
        with pytest.raises(ParseError, match=r"Legacy .doc"):
            parse_resume(b"\xd0\xcf\x11\xe0" + b"x" * 200, "resume.doc")

    def test_unsupported_extension(self):
        with pytest.raises(ParseError, match="Unsupported"):
            parse_resume(b"hello", "resume.png")

    def test_extension_spoofing_detected(self):
        with pytest.raises(ParseError, match="isn't a valid PDF"):
            parse_resume(b"not a pdf" * 20, "resume.pdf")

    def test_too_little_text(self):
        with pytest.raises(ParseError, match="Almost no text"):
            parse_resume(b"Just a name", "resume.txt")


class TestChunking:
    def test_split_respects_budget_and_overlaps(self):
        text = " ".join(f"Sentence number {i} describes an accomplishment in some detail." for i in range(40))
        chunks = split_text(text, max_chars=300)
        assert len(chunks) > 3
        assert all(len(c) <= 300 for c in chunks)
        # one sentence of overlap between neighbours
        last_sentence = chunks[0].split(". ")[-1].rstrip(".")
        assert last_sentence in chunks[1]

    def test_hard_wraps_run_on_text(self):
        chunks = split_text("x" * 2000, max_chars=500)
        assert all(len(c) <= 500 for c in chunks)
        assert sum(len(c) for c in chunks) >= 2000

    def test_resume_chunks_are_structured(self):
        raw = read_sample("resume_ada_strong.txt")
        resume = mock_resume(raw)
        chunks = build_resume_chunks(resume, raw)
        sections = {c.section for c in chunks}
        assert {"experience", "skills", "education"} <= sections
        experience = [c for c in chunks if c.section == "experience"]
        assert experience[0].text.startswith("Senior Software Engineer at Streamline Analytics")
        assert set(build_section_texts(resume)) >= {"skills", "experience", "education"}
